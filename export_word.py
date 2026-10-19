"""
Script tự động trích xuất toàn bộ hồ sơ sửa lỗi và mã nguồn hoàn chỉnh ra file Microsoft Word (.docx)
Chạy lệnh: python export_word.py
"""
import os
import sys

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_word_report(output_filename="HO_SO_SUA_LOI_VA_MA_NGUON_DATN_20224083.docx"):
    doc = Document()

    # Tiêu đề chính
    title = doc.add_heading("HỒ SƠ RÀ SOÁT, SỬA LỖI & MÃ NGUỒN HOÀN CHỈNH ĐATN", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_info = doc.add_paragraph()
    p_info.add_run("Đề tài: ").bold = True
    p_info.add_run("Hệ thống giám sát và phát hiện biển quảng cáo/biển hiệu lấn chiếm vỉa hè qua camera CCTV cố định\n")
    p_info.add_run("Sinh viên: ").bold = True
    p_info.add_run("Nguyễn Đức Tài Năng (MSSV: 20224083) - Đại học Bách Khoa Hà Nội")

    files_to_export = [
        (
            "1. Module Xác Minh Chuỗi Thời Gian (tracking/verifier.py)",
            "tracking/verifier.py",
            "Đưa stationary_window_seconds (1.0s) và spatial_cooldown_radius (50.0px) vào tham số khởi tạo __init__ đồng bộ với settings.yaml. "
            "Bổ sung cơ chế Perspective-Adaptive Stationary Check chuẩn hóa ngưỡng dịch chuyển theo kích thước đường chéo Bounding Box để chống biến dạng phối cảnh ở cự ly xa.",
        ),
        (
            "2. Script Đánh Giá Toàn Hệ Thống (evaluation/evaluate_videos.py)",
            "evaluation/evaluate_videos.py",
            "Sửa lỗi lệch tên khóa cấu hình YAML ('detector', 'spatial', 'temporal'). "
            "Bổ sung kiểm chứng lân cận không gian cho Pass 1 (direct_id), tính Spatial IoU cực đại trên toàn chuỗi lịch sử boxes kèm tâm trung vị, "
            "áp dụng thuật toán Hungarian (scipy.optimize.linear_sum_assignment) cho Pass 2 và gọi roi_manager.validate_resolution(width, height).",
        ),
        (
            "3. Điểm Vào Điều Hành Trung Tâm (main.py)",
            "main.py",
            "Đặt default=None cho các cờ --conf, --overlap, --confirm-frames, --cooldown trong argparse để ưu tiên đọc cấu hình chuẩn từ configs/settings.yaml. "
            "Đồng bộ confirm_seconds = 3.0s và truyền đầy đủ spatial_cooldown_radius vào TemporalVerifier.",
        ),
        (
            "4. File Cấu Hình Hệ Thống Tập Trung (configs/settings.yaml)",
            "configs/settings.yaml",
            "Giữ nguyên cấu hình chuẩn hóa v8.0 đã đóng băng phục vụ đối sánh Single Source of Truth.",
        ),
    ]

    for heading_text, filepath, reason_text in files_to_export:
        doc.add_heading(heading_text, level=1)
        p_reason = doc.add_paragraph()
        r_lbl = p_reason.add_run("Lý do & Giải pháp nâng cấp: ")
        r_lbl.bold = True
        r_lbl.font.color.rgb = RGBColor(0, 102, 204)
        p_reason.add_run(reason_text)

        doc.add_heading(f"Mã nguồn tệp: {filepath}", level=2)
        if os.path.exists(filepath):
            with open(filepath, "r", encoding="utf-8") as f:
                code_content = f.read()
        else:
            code_content = f"# [Không tìm thấy tệp {filepath} tại đường dẫn hiện tại]"

        code_para = doc.add_paragraph()
        code_run = code_para.add_run(code_content)
        code_run.font.name = "Consolas"
        code_run.font.size = Pt(8.5)

    doc.save(output_filename)
    print(f"[THÀNH CÔNG] Đã xuất file Word hoàn chỉnh tại: {os.path.abspath(output_filename)}")

if __name__ == "__main__":
    create_word_report()
