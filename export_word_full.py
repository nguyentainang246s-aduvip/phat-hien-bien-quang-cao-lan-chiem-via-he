"""
Script tự động xuất Hồ sơ Rà soát, Lý do sửa, Cách sửa và Toàn bộ Mã nguồn (Nhóm 1, 2, 3) ra file Word (.docx)
Chạy lệnh: python export_word_full.py
"""
import os
import sys

if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


def generate_full_thesis_code_docx(output_path="HO_SO_MA_NGUON_VA_SUA_LOI_TOAN_DIEN_DATN_20224083.docx"):
    doc = Document()

    # Căn lề chuẩn báo cáo
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(0.8)

    # Tiêu đề chính
    title = doc.add_heading("HỒ SƠ RÀ SOÁT, NÂNG CẤP & TỔNG HỢP MÃ NGUỒN ĐATN", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = doc.add_paragraph()
    meta.add_run("Đề tài: ").bold = True
    meta.add_run("Hệ thống giám sát và phát hiện biển quảng cáo/biển hiệu lấn chiếm vỉa hè qua camera CCTV cố định\n")
    meta.add_run("Sinh viên thực hiện: ").bold = True
    meta.add_run("Nguyễn Đức Tài Năng — MSSV: 20224083 (Đại học Bách Khoa Hà Nội)\n")
    meta.add_run("Phạm vi tài liệu: ").bold = True
    meta.add_run("Bao gồm Lý do sửa, Cách khắc phục và Toàn bộ mã nguồn Nhóm 1 (Bắt buộc), Nhóm 2 (Quan trọng), Nhóm 3 (Bổ sung).")

    items = [
        (
            "NHÓM 1 - FILE 1: violation/checker.py (Quyết định Vi phạm Hình học Không gian)",
            "violation/checker.py",
            "Chưa đưa tham số ground_contact_rule và base_margin vào hàm khởi tạo __init__ để đồng bộ với settings.yaml.",
            "Bổ sung ground_contact_rule: bool = True và base_margin: float = 15.0 vào __init__ của ViolationChecker.",
        ),
        (
            "NHÓM 1 - FILE 2: tracking/verifier.py (Bộ xác minh Chuỗi thời gian & Máy trạng thái v2.2)",
            "tracking/verifier.py",
            "Ngưỡng dịch chuyển cố định 50px chưa thích ứng theo chiều sâu phối cảnh; thiếu tham số spatial_cooldown_radius ở __init__.",
            "Bổ sung cơ chế Perspective-Adaptive Stationary Check chuẩn hóa theo đường chéo BBox và đồng bộ tham số với settings.yaml.",
        ),
        (
            "NHÓM 1 - FILE 3: evaluation/evaluate_videos.py (Đánh giá Toàn hệ thống & Hungarian Matching)",
            "evaluation/evaluate_videos.py",
            "Lệch tên khóa YAML ('detection' vs 'detector'); Pass 1 chưa kiểm tra khoảng cách không gian; Pass 2 dùng Greedy thay vì Hungarian.",
            "Đọc đúng khóa YAML chuẩn; bổ sung _compute_best_spatial_iou xét lịch sử boxes; áp dụng scipy.optimize.linear_sum_assignment và validate_resolution.",
        ),
        (
            "NHÓM 1 - FILE 4: tracking/tracker.py (Bộ bám vết đa đối tượng ByteTrack)",
            "tracking/tracker.py",
            "Cần đảm bảo giải phóng và reset bộ lọc Kalman an toàn khi chuyển đổi nguồn video hoặc tua lại luồng.",
            "Duy trì persist=True trong quá trình theo dõi liên tục và cung cấp phương thức reset() làm sạch trạng thái predictor.trackers.",
        ),
        (
            "NHÓM 2 - FILE 5: camera/stream_reader.py (Đọc luồng RTSP Zero-Latency & Auto-Reconnect)",
            "camera/stream_reader.py",
            "Nguy cơ xung đột đa luồng (Race Condition) khi _connect() giải phóng self.cap lúc luồng _capture_worker cũ chưa kết thúc.",
            "Bổ sung hàm _stop_worker_thread() gọi _thread.join(timeout=1.5) trước khi đóng/mở lại kết nối cv2.VideoCapture.",
        ),
        (
            "NHÓM 2 - FILE 6: evidence/saver.py (Trích xuất Bằng chứng Pháp lý & Bảo vệ Quyền riêng tư)",
            "evidence/saver.py",
            "Khi truyền clean_frame vào save_evidence(), ảnh toàn cảnh chưa có khung đỏ chỉ rõ biển hiệu nào vi phạm.",
            "Cắt ảnh cận cảnh (Crop) từ clean_frame trước để giữ ảnh Crop sạch 100%, sau đó vẽ khung đỏ kèm mã ID và % lấn chiếm lên ảnh toàn cảnh.",
        ),
        (
            "NHÓM 2 - FILE 7: roi/roi_manager.py (Quản lý Đa giác Vỉa hè Multi-ROI & Tự động Co giãn)",
            "roi/roi_manager.py",
            "Đảm bảo tính chính xác hình học khi độ phân giải video thực tế khác với độ phân giải lúc vẽ cấu hình ROI.",
            "Thực thi validate_resolution() và adapt_to_resolution() dựa trên base_polygons gốc trước mọi phép tính không gian.",
        ),
        (
            "NHÓM 2 - FILE 8: database/db.py (Cơ sở dữ liệu SQLite WAL Mode & Khử trùng lặp)",
            "database/db.py",
            "Đồng bộ schema bảng violations hỗ trợ quy trình duyệt Human-in-the-Loop và chống ghi trùng khi khởi động lại.",
            "Truy vấn tham số hóa chống SQL Injection, bật WAL Mode và kiểm tra trùng lặp tọa độ BBox trong cửa sổ 5 phút.",
        ),
        (
            "NHÓM 2 - FILE 9: requirements.txt (Danh mục Thư viện Chuẩn hóa)",
            "requirements.txt",
            "Thiếu khai báo onnxruntime và scipy trong danh sách phụ thuộc.",
            "Bổ sung onnxruntime>=1.16.0, scipy>=1.10.0, psutil>=5.9.0 để đảm bảo tái lập môi trường thực nghiệm 100%.",
        ),
        (
            "NHÓM 3 - FILE 10: main.py (Điểm vào Điều hành Hệ thống Trung tâm)",
            "main.py",
            "Các cờ CLI --conf và --confirm-frames đặt giá trị mặc định cứng làm ghi đè cấu hình trong settings.yaml.",
            "Đặt default=None trong argparse cho các tham số ngưỡng để ưu tiên nạp từ configs/settings.yaml.",
        ),
        (
            "NHÓM 3 - FILE 11: dashboard/web_app.py (Web Dashboard Human-in-the-Loop & ROI Studio)",
            "dashboard/web_app.py",
            "Đồng bộ khóa luồng khi suy luận AI trực tuyến, gọi tracker.reset() khi tua video và đồng bộ confirm_seconds = 3.0s.",
            "Sử dụng _TRACKER_INFER_LOCK bảo vệ luồng MJPEG, truyền simulated timestamp chuẩn và lọc ngày linh hoạt.",
        ),
        (
            "NHÓM 3 - FILE 12: scripts/train_colab.py (Kịch bản Huấn luyện YOLOv8n trên Google Colab)",
            "scripts/train_colab.py",
            "Khóa Roboflow API Key đang được viết cứng (hardcoded) trong mã nguồn.",
            "Đọc khóa từ biến môi trường os.environ.get('ROBOFLOW_API_KEY') để đảm bảo an toàn bảo mật mã nguồn.",
        ),
    ]

    for title_str, rel_path, reason_str, fix_str in items:
        doc.add_heading(title_str, level=1)

        p_r = doc.add_paragraph()
        r1 = p_r.add_run("• Lý do cần sửa / nâng cấp: ")
        r1.bold = True
        r1.font.color.rgb = RGBColor(192, 0, 0)
        p_r.add_run(reason_str)

        p_f = doc.add_paragraph()
        r2 = p_f.add_run("• Cách sửa & Giải pháp kỹ thuật: ")
        r2.bold = True
        r2.font.color.rgb = RGBColor(0, 112, 192)
        p_f.add_run(fix_str)

        doc.add_heading(f"Mã nguồn tệp: {rel_path}", level=2)
        if os.path.exists(rel_path):
            with open(rel_path, "r", encoding="utf-8") as f:
                content = f.read()
        else:
            content = f"# [Không tìm thấy tệp {rel_path} tại thư mục hiện hành]"

        p_code = doc.add_paragraph()
        run_code = p_code.add_run(content)
        run_code.font.name = "Consolas"
        run_code.font.size = Pt(8.0)

    doc.save(output_path)
    print(f"[THÀNH CÔNG] Đã xuất toàn bộ hồ sơ mã nguồn ra file Word: {os.path.abspath(output_path)}")


if __name__ == "__main__":
    generate_full_thesis_code_docx()
