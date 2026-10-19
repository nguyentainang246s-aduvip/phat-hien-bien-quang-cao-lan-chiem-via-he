"""
Script xuất bản bài học Bài 3 thành file Word .docx chuyên nghiệp
Lưu vào thư mục: tai_lieu_hoc_tap/Bai_03_Mo_Hinh_AI_YOLOv8_va_Bounding_Box.docx
"""

import os
import sys
from pathlib import Path

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

ROOT_DIR = Path(__file__).resolve().parent.parent
OUTPUT_DIR = ROOT_DIR / "tai_lieu_hoc_tap"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
OUTPUT_FILE = str(OUTPUT_DIR / "Bai_03_Mo_Hinh_AI_YOLOv8_va_Bounding_Box.docx")

COLOR_PRIMARY = RGBColor(30, 58, 138)     # Navy #1E3A8A
COLOR_SECONDARY = RGBColor(15, 118, 110)  # Teal #0F766E
COLOR_DARK = RGBColor(30, 41, 59)         # Slate Dark #1E293B
COLOR_MUTED = RGBColor(100, 116, 139)     # Slate Gray #64748B
COLOR_RED = RGBColor(220, 38, 38)         # Red

HEX_PRIMARY = "1E3A8A"
HEX_SECONDARY = "0F766E"
HEX_BG_LIGHT = "F8FAFC"
HEX_BG_CODE = "F1F5F9"
HEX_BORDER = "CBD5E1"
HEX_HEADER_BG = "1E3A8A"
HEX_CALLOUT_BG = "EFF6FF"
HEX_CALLOUT_BORDER = "3B82F6"
HEX_WARNING_BG = "FEF2F2"
HEX_WARNING_BORDER = "EF4444"
HEX_CONFIRMED = "DBEAFE"


def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def set_cell_background(cell, hex_color):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shd)


def set_table_borders(table, hex_color=HEX_BORDER):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="{hex_color}"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{hex_color}"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def add_callout_box(doc, text, title="LƯU Ý QUAN TRỌNG", is_warning=False):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    
    bg_color = HEX_WARNING_BG if is_warning else HEX_CALLOUT_BG
    border_color = HEX_WARNING_BORDER if is_warning else HEX_CALLOUT_BORDER
    
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'  <w:top w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    cell._tc.get_or_add_tcPr().append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run_title = p.add_run(f"{title}: ")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(10.5)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(220, 38, 38) if is_warning else COLOR_PRIMARY
    
    run_text = p.add_run(text)
    run_text.font.name = "Arial"
    run_text.font.size = Pt(10)
    run_text.font.color.rgb = COLOR_DARK
    
    doc.add_paragraph()


def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    
    set_cell_background(cell, HEX_BG_CODE)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="16" w:space="0" w:color="{HEX_PRIMARY}"/>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'</w:tcBorders>'
    )
    cell._tc.get_or_add_tcPr().append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(13)
    
    run = p.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(30, 41, 59)
    
    doc.add_paragraph()


def build_document():
    doc = Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    # --- HEADER TIÊU ĐỀ ---
    p_header = doc.add_paragraph()
    p_header.paragraph_format.space_after = Pt(2)
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_uni = p_header.add_run("ĐẠI HỌC BÁCH KHOA HÀ NỘI — TRƯỜNG CÔNG NGHỆ THÔNG TIN VÀ TRUYỀN THÔNG\n")
    r_uni.font.name = "Arial"
    r_uni.font.size = Pt(10)
    r_uni.font.bold = True
    r_uni.font.color.rgb = COLOR_MUTED
    
    r_sub = p_header.add_run("ĐỒ ÁN TỐT NGHIỆP: HỆ THỐNG PHÁT HIỆN BIỂN QUẢNG CÁO LẤN CHIẾM VỈA HÈ QUA CCTV\n")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(11)
    r_sub.font.bold = True
    r_sub.font.color.rgb = COLOR_SECONDARY
    
    r_badge = p_header.add_run("CHƯƠNG TRÌNH ĐÀO TẠO NỘI BỘ — HỌC HIỂU ĐỒ ÁN TỪ CON SỐ 0")
    r_badge.font.name = "Arial"
    r_badge.font.size = Pt(9.5)
    r_badge.font.italic = True
    r_badge.font.color.rgb = COLOR_MUTED

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(4)
    p_div.paragraph_format.space_after = Pt(14)
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_div = p_div.add_run("━" * 50)
    r_div.font.color.rgb = COLOR_PRIMARY

    # Tiêu đề bài học
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_after = Pt(8)
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_b3 = p_title.add_run("BÀI 3: MÔ HÌNH AI YOLOV8N VÀ BOUNDING BOX\n")
    r_b3.font.name = "Arial"
    r_b3.font.size = Pt(18)
    r_b3.font.bold = True
    r_b3.font.color.rgb = COLOR_PRIMARY
    
    r_b3_sub = p_title.add_run("Cách mạng Deep Learning tìm kiếm và khoanh vùng biển quảng cáo trong ảnh")
    r_b3_sub.font.name = "Arial"
    r_b3_sub.font.size = Pt(13)
    r_b3_sub.font.bold = True
    r_b3_sub.font.color.rgb = COLOR_SECONDARY

    # Bảng siêu dữ liệu
    tbl_meta = doc.add_table(rows=3, cols=2)
    tbl_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Tệp mã nguồn tương ứng:", "detection/detector.py (Tham chiếu: scripts/demos/demo_yolo.py, models/best.onnx)"),
        ("Cấp độ bài học:", "Giai đoạn 2 — Deep Learning & Object Detection"),
        ("Đối tượng học tập:", "Người học từ con số 0, cần hiểu cách YOLO tạo ra khung chữ nhật và độ tin cậy")
    ]
    for i, (k, v) in enumerate(meta_data):
        c0 = tbl_meta.cell(i, 0)
        c1 = tbl_meta.cell(i, 1)
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        set_cell_background(c0, HEX_BG_LIGHT)
        set_cell_background(c1, HEX_BG_LIGHT)
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        r0.font.name = "Arial"
        r0.font.size = Pt(9.5)
        r0.font.bold = True
        r0.font.color.rgb = COLOR_DARK
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v)
        r1.font.name = "Arial"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = COLOR_DARK
    set_table_borders(tbl_meta, HEX_BORDER)
    doc.add_paragraph()

    # --- PHẦN 1: HÔM NAY HỌC GÌ? ---
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(4)
    r_h1 = h1.add_run("PHẦN 1: HÔM NAY HỌC GÌ?")
    r_h1.font.name = "Arial"
    r_h1.font.size = Pt(13)
    r_h1.font.bold = True
    r_h1.font.color.rgb = COLOR_PRIMARY

    p_p1 = doc.add_paragraph()
    p_p1.paragraph_format.line_spacing = Pt(14)
    p_p1.paragraph_format.space_after = Pt(6)
    p_p1.add_run(
        "• Tên bài học: Phát hiện vật thể (Object Detection) với mô hình YOLOv8n và bản chất của Bounding Box.\n"
        "• Mục tiêu cốt lõi:\n"
        "   1. Hiểu Object Detection là gì: vừa xác định vị trí (Where) vừa phân loại vật thể (What).\n"
        "   2. Hiểu cấu trúc Bounding Box [x1, y1, x2, y2] và Độ tự tin (Confidence Score).\n"
        "   3. Đọc hiểu toàn bộ module detection/detector.py: cách nạp mô hình best.pt/best.onnx và lấy kết quả dự đoán.\n"
        "• Vì sao quan trọng đối với đồ án?\n"
        "   Mô hình AI là 'đôi mắt' của hệ thống. Nếu AI không phát hiện ra biển quảng cáo hoặc đưa ra tọa độ khung bao sai lệch, thì tầng phán quyết hình học và lưu bằng chứng phía sau sẽ không có dữ liệu để hoạt động."
    )

    # --- PHẦN 2: KIẾN THỨC NỀN TẢNG ---
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(4)
    r_h2 = h2.add_run("PHẦN 2: KIẾN THỨC NỀN TẢNG")
    r_h2.font.name = "Arial"
    r_h2.font.size = Pt(13)
    r_h2.font.bold = True
    r_h2.font.color.rgb = COLOR_PRIMARY

    p_p2 = doc.add_paragraph()
    p_p2.paragraph_format.line_spacing = Pt(14)
    p_p2.paragraph_format.space_after = Pt(6)
    p_p2.add_run(
        "1. Phát hiện vật thể (Object Detection) là gì?\n"
        "Khác với Phân loại ảnh (Image Classification - chỉ trả lời trong ảnh có biển hiệu hay không), Object Detection giải quyết 2 bài toán cùng lúc:\n"
        "   - Định vị (Localization): Vẽ một khung chữ nhật bao quanh vật thể.\n"
        "   - Phân loại (Classification): Xác định vật thể trong khung chữ nhật đó thuộc nhóm nào (trong đồ án của chúng ta là class: bien_quang_cao).\n\n"
        "2. Khung bao vật thể (Bounding Box - viết tắt là BBox):\n"
        "Là một hình chữ nhật đóng khung đối tượng trên ảnh. Trong OpenCV và YOLOv8, chuẩn xyxy được định nghĩa bởi 4 con số nguyên:\n"
        "   - (x1, y1): Tọa độ góc trên cùng bên trái của khung bao.\n"
        "   - (x2, y2): Tọa độ góc dưới cùng bên phải của khung bao.\n"
        "   - Từ 4 con số này, ta suy ra ngay đáy của biển hiệu nằm ở đường ngang y = y2, và điểm giữa chân biển có tọa độ x = (x1 + x2) / 2!\n\n"
        "3. Độ tin cậy (Confidence Score - Conf):\n"
        "Là một số thực nằm trong khoảng từ 0.0 đến 1.0 (hoặc 0% đến 100%), đại diện cho mức độ 'chắc chắn' của mô hình AI rằng trong ô đó có chứa biển hiệu.\n"
        "   - Ví dụ: Conf = 0.85 nghĩa là AI chắc chắn 85% đây là biển quảng cáo.\n"
        "   - Ngưỡng lọc (Conf Threshold): Chỉ giữ lại các dự đoán có Conf >= ngưỡng (ví dụ: >= 0.25 hoặc 0.35), loại bỏ các phỏng đoán mơ hồ.\n\n"
        "4. Mô hình YOLOv8n (Nano):\n"
        "YOLO là chữ viết tắt của 'You Only Look Once' (Chỉ cần nhìn một lần). Phiên bản Nano ('n') là phiên bản nhỏ gọn nhất (chỉ 3.0 triệu tham số, dung lượng ~6.4 MB), được thiết kế để chạy mượt mà ngay trên CPU máy tính thông thường mà không cần card màn hình rời (GPU) đắt tiền."
    )

    add_callout_box(
        doc,
        "Đồ án huấn luyện YOLO với 1 class duy nhất là 'bien_quang_cao'. Mô hình AI CHỈ LÀM NHIỆM VỤ TÌM BIỂN, "
        "không kết luận vi phạm hay không. Tầng thuật toán hình học phía sau mới quyết định biển có lấn vỉa hè hay không!",
        "NGUYÊN TẮC PHÂN TÁCH TRÁCH NHIỆM",
        is_warning=False
    )

    # --- PHẦN 3: VÍ DỤ ĐỜI SỐNG ---
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(12)
    h3.paragraph_format.space_after = Pt(4)
    r_h3 = h3.add_run("PHẦN 3: VÍ DỤ ĐỜI SỐNG MINH HỌA")
    r_h3.font.name = "Arial"
    r_h3.font.size = Pt(13)
    r_h3.font.bold = True
    r_h3.font.color.rgb = COLOR_PRIMARY

    p_p3 = doc.add_paragraph()
    p_p3.paragraph_format.line_spacing = Pt(14)
    p_p3.paragraph_format.space_after = Pt(6)
    p_p3.add_run(
        "Hãy tưởng tượng bạn thuê một bác bảo vệ đứng quan sát camera trước cổng trường:\n"
        "• Bạn đưa cho bác bảo vệ 1000 tấm ảnh chụp các loại biển hiệu (standee, biển hộp đèn, biển chân sắt) để bác học cách nhận biết.\n"
        "• Khi camera chiếu một khung hình đường phố mới, bác bảo vệ lập tức chỉ tay vào màn hình: 'Cháu ơi, ở ô chữ nhật góc dưới bên phải này, bác chắc chắn 90% là có một tấm biển quảng cáo!'.\n"
        "• Trong đồ án của chúng ta: Mô hình YOLOv8n chính là bác bảo vệ tinh mắt đó. Mỗi khung hình frame đưa vào, YOLO sẽ 'chỉ tay' và trả về danh sách các ô chữ nhật chứa biển hiệu kèm mức độ tin cậy của nó!"
    )

    # --- PHẦN 4: LIÊN HỆ VỚI CODE THẬT ---
    h4 = doc.add_paragraph()
    h4.paragraph_format.space_before = Pt(12)
    h4.paragraph_format.space_after = Pt(4)
    r_h4 = h4.add_run("PHẦN 4: LIÊN HỆ VỚI MÃ NGUỒN THỰC TẾ (detection/detector.py)")
    r_h4.font.name = "Arial"
    r_h4.font.size = Pt(13)
    r_h4.font.bold = True
    r_h4.font.color.rgb = COLOR_PRIMARY

    p_p4 = doc.add_paragraph()
    p_p4.paragraph_format.line_spacing = Pt(14)
    p_p4.paragraph_format.space_after = Pt(4)
    p_p4.add_run("Mã nguồn bọc mô hình AI trong tệp detection/detector.py:")

    code_content = (
        'from ultralytics import YOLO\n\n'
        'class YOLODetector:\n'
        '    def __init__(self, model_path: str = "models/best.pt", conf_threshold: float = 0.35, imgsz: int = 1280):\n'
        '        self.model_path = model_path\n'
        '        self.conf_threshold = conf_threshold\n'
        '        self.imgsz = imgsz\n'
        '        # Nạp trọng số mô hình đã huấn luyện\n'
        '        self.model = YOLO(model_path)\n'
        '        self.class_names = self.model.names\n\n'
        '    def detect(self, frame, imgsz: int = None) -> list:\n'
        '        infer_imgsz = imgsz if imgsz is not None else self.imgsz\n'
        '        # 1. Chạy mạng nơ-ron dự đoán trên frame ảnh\n'
        '        results = self.model.predict(frame, conf=self.conf_threshold, imgsz=infer_imgsz, verbose=False)\n'
        '        detections = []\n\n'
        '        if len(results) == 0 or results[0].boxes is None:\n'
        '            return detections\n\n'
        '        # 2. Bóc tách từng Bounding Box được tìm thấy\n'
        '        boxes = results[0].boxes\n'
        '        for box in boxes:\n'
        '            # Lấy tọa độ (x1, y1, x2, y2), độ tự tin (conf), mã lớp (cls_id)\n'
        '            xyxy = box.xyxy[0].cpu().numpy().astype(int)\n'
        '            conf = float(box.conf[0].cpu().numpy())\n'
        '            cls_id = int(box.cls[0].cpu().numpy())\n'
        '            cls_name = self.class_names.get(cls_id, f"class_{cls_id}")\n\n'
        '            detections.append({\n'
        '                "box": (int(xyxy[0]), int(xyxy[1]), int(xyxy[2]), int(xyxy[3])),\n'
        '                "conf": round(conf, 2),\n'
        '                "class_id": cls_id,\n'
        '                "class_name": cls_name\n'
        '            })\n'
        '        return detections'
    )
    add_code_block(doc, code_content)

    # --- PHẦN 5: GIẢI THÍCH CODE TỪNG DÒNG ---
    h5 = doc.add_paragraph()
    h5.paragraph_format.space_before = Pt(12)
    h5.paragraph_format.space_after = Pt(4)
    r_h5 = h5.add_run("PHẦN 5: GIẢI THÍCH CODE TỪNG DÒNG VÀ Ý NGHĨA KỸ THUẬT")
    r_h5.font.name = "Arial"
    r_h5.font.size = Pt(13)
    r_h5.font.bold = True
    r_h5.font.color.rgb = COLOR_PRIMARY

    p_p5 = doc.add_paragraph()
    p_p5.paragraph_format.line_spacing = Pt(14)
    p_p5.paragraph_format.space_after = Pt(6)
    p_p5.add_run(
        "1. Khởi tạo YOLODetector (__init__):\n"
        "• self.model = YOLO(model_path): Nạp file trọng số best.pt (kết quả sau 92 epochs huấn luyện). Trọng số này chứa các tham số ma trận đã được tối ưu để nhận biết các đặc trưng hình học của biển quảng cáo.\n"
        "• conf_threshold = 0.35: Thiết lập màng lọc. Mọi vật thể có độ tự tin dưới 35% sẽ bị loại bỏ ngay để tránh báo động giả từ các vật thể na ná như thùng rác, cửa sổ.\n"
        "• imgsz = 1280: Kích thước độ phân giải đưa vào mạng AI. Vì camera CCTV đặt trên cao góc rộng, các biển quảng cáo ở xa thường nhỏ (Small Objects), nên cần đẩy độ phân giải lên 1280 để mạng nhìn rõ chi tiết.\n\n"
        "2. Quá trình dự đoán (Hàm detect):\n"
        "• results = self.model.predict(frame, conf=..., imgsz=..., verbose=False): ĐÂY LÀ DÒNG CODE SUY LUẬN AI (INFERENCE). Bức ảnh frame từ camera được đưa qua hàng chục tầng nơ-ron tích chập (Convolutional Layers) để tìm vật thể. Tham số verbose=False giúp màn hình console không bị in tràn lan hàng trăm dòng log mỗi giây.\n"
        "• box.xyxy[0].cpu().numpy().astype(int): Đầu ra của mạng nơ-ron là một PyTorch Tensor nằm trên bộ nhớ tính toán. Lệnh này chuyển dữ liệu Tensor về mảng số nguyên thông thường của Python: [x1, y1, x2, y2].\n"
        "• box.conf[0]: Trích xuất độ tự tin (ví dụ: 0.88).\n"
        "• detections.append({...}): Đóng gói thông tin của từng tấm biển thành một từ điển (Dictionary) rõ ràng, sạch sẽ, chuẩn bị nạp sang tầng tiếp theo."
    )

    # --- PHẦN 6: MINH HỌA QUÁ TRÌNH CHẠY ---
    h6 = doc.add_paragraph()
    h6.paragraph_format.space_before = Pt(12)
    h6.paragraph_format.space_after = Pt(4)
    r_h6 = h6.add_run("PHẦN 6: MINH HỌA QUÁ TRÌNH DỮ LIỆU BIẾN ĐỔI QUA MẠNG AI")
    r_h6.font.name = "Arial"
    r_h6.font.size = Pt(13)
    r_h6.font.bold = True
    r_h6.font.color.rgb = COLOR_PRIMARY

    p_p6 = doc.add_paragraph()
    p_p6.paragraph_format.space_after = Pt(6)
    p_p6.add_run("Bảng mô phỏng dữ liệu đầu vào và kết quả đầu ra của hàm detector.detect(frame):")

    tbl_flow = doc.add_table(rows=4, cols=4)
    tbl_flow.alignment = WD_TABLE_ALIGNMENT.CENTER
    flow_headers = ["Đầu vào (Input)", "Xử lý mạng YOLOv8n", "Kết quả đầu ra (Detections)", "Ý nghĩa thực tế"]
    for j, h in enumerate(flow_headers):
        c = tbl_flow.cell(0, j)
        set_cell_background(c, HEX_HEADER_BG)
        set_cell_margins(c, 100, 100, 120, 120)
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    flow_rows = [
        ("Frame ảnh 720p\n(720, 1280, 3)", "Phát hiện 1 biển standee đứng\nConf = 0.87", "[{'box': (410, 350, 490, 520),\n'conf': 0.87, 'class_name': 'bien_quang_cao'}]", "Tìm thấy 1 biển hiệu tại tọa độ x1=410, y1=350, x2=490, y2=520"),
        ("Frame ảnh 720p\n(720, 1280, 3)", "Phát hiện 1 biển vẫy gắn tường\nConf = 0.76", "[{'box': (120, 100, 200, 250),\n'conf': 0.76, 'class_name': 'bien_quang_cao'}]", "Tìm thấy 1 biển treo ở góc trên bên trái"),
        ("Frame chỉ có người\n& xe máy đi lại", "Không có đối tượng nào đạt\nConf >= 0.35", "[] (Danh sách rỗng)", "Không có biển quảng cáo nào trong khung hình")
    ]
    for i, row in enumerate(flow_rows):
        for j, val in enumerate(row):
            c = tbl_flow.cell(i + 1, j)
            bg = HEX_BG_LIGHT if i % 2 == 0 else "FFFFFF"
            set_cell_background(c, bg)
            set_cell_margins(c, 80, 80, 100, 100)
            p = c.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Arial"
            r.font.size = Pt(9)
            r.font.color.rgb = COLOR_DARK
    set_table_borders(tbl_flow, HEX_BORDER)
    doc.add_paragraph()

    # --- PHẦN 7: LỖI THƯỜNG GẶP ---
    h7 = doc.add_paragraph()
    h7.paragraph_format.space_before = Pt(12)
    h7.paragraph_format.space_after = Pt(4)
    r_h7 = h7.add_run("PHẦN 7: LỖI THƯỜNG GẶP CỦA NGƯỜI MỚI BẮT ĐẦU VÀ CÁCH KHẮC PHỤC")
    r_h7.font.name = "Arial"
    r_h7.font.size = Pt(13)
    r_h7.font.bold = True
    r_h7.font.color.rgb = COLOR_PRIMARY

    p_p7 = doc.add_paragraph()
    p_p7.paragraph_format.line_spacing = Pt(14)
    p_p7.paragraph_format.space_after = Pt(6)
    p_p7.add_run(
        "1. Lỗi 'CUDA out of memory' hoặc chạy quá chậm trên máy không có GPU:\n"
        "• Nguyên nhân: Cố ép PyTorch chạy trên GPU khi card đồ họa không đủ VRAM.\n"
        "• Khắc phục: Đồ án đã được tối ưu hóa sẵn mô hình Nano (YOLOv8n) và hỗ trợ xuất sang ONNX Runtime (models/best.onnx) chạy cực kỳ mượt mà trên CPU với thông lượng 25+ FPS.\n\n"
        "2. Nhầm lẫn định dạng tọa độ Box:\n"
        "• Nhầm lẫn: Tưởng [x1, y1, x2, y2] là [x, y, width, height].\n"
        "• Thực tế: xyxy là tọa độ 2 điểm chéo nhau: Góc trên-trái (x1, y1) và Góc dưới-phải (x2, y2). Chiều rộng Width thực chất là x2 - x1, còn chiều cao Height là y2 - y1."
    )

    # --- PHẦN 8: BÀI TẬP KIỂM TRA ---
    h8 = doc.add_paragraph()
    h8.paragraph_format.space_before = Pt(12)
    h8.paragraph_format.space_after = Pt(4)
    r_h8 = h8.add_run("PHẦN 8: BÀI TẬP KIỂM TRA HIỂU BÀI (TỰ ÔN TẬP)")
    r_h8.font.name = "Arial"
    r_h8.font.size = Pt(13)
    r_h8.font.bold = True
    r_h8.font.color.rgb = COLOR_PRIMARY

    p_p8 = doc.add_paragraph()
    p_p8.paragraph_format.line_spacing = Pt(14)
    p_p8.paragraph_format.space_after = Pt(10)
    p_p8.add_run(
        "Bạn hãy tự trả lời 3 câu hỏi dưới đây trước khi chuyển sang Bài học số 4:\n\n"
        "• Câu 1: Giả sử mô hình YOLO trả về Bounding Box của một tấm biển là (x1=200, y1=100, x2=300, y2=500). Hỏi đáy tiếp xúc mặt đất của tấm biển này nằm ở đường thẳng nằm ngang y bằng bao nhiêu (y=100 hay y=500)?\n\n"
        "• Câu 2: Biến 'conf' (Confidence Score) đại diện cho điều gì? Nếu bạn đặt ngưỡng conf_threshold = 0.50, điều gì sẽ xảy ra với các biển hiệu ở xa có conf = 0.42?\n\n"
        "• Câu 3: Tại sao trong đồ án này, chúng ta chỉ huấn luyện YOLO nhận diện class 'bien_quang_cao', mà KHÔNG huấn luyện nhận diện class 'bien_lan_chiem_via_he'?"
    )

    doc.save(OUTPUT_FILE)
    print(f"[THÀNH CÔNG] Đã tạo tệp tài liệu bài học: {OUTPUT_FILE}")


if __name__ == "__main__":
    build_document()
