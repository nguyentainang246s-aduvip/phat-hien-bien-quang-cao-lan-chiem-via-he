"""
Script xuất bản bài học Bài 7 thành file Word .docx chuyên nghiệp
Lưu vào thư mục: tai_lieu_hoc_tap/Bai_07_May_Trang_Thai_va_Bo_Loc_Thoi_Gian.docx
"""

import os
import sys
from pathlib import Path

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

ROOT_DIR = Path(__file__).resolve().parent.parent
OUTPUT_DIR = ROOT_DIR / "tai_lieu_hoc_tap"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
OUTPUT_FILE = str(OUTPUT_DIR / "Bai_07_May_Trang_Thai_va_Bo_Loc_Thoi_Gian.docx")

COLOR_PRIMARY = RGBColor(30, 58, 138)     # Navy #1E3A8A
COLOR_SECONDARY = RGBColor(15, 118, 110)  # Teal #0F766E
COLOR_DARK = RGBColor(30, 41, 59)         # Slate Dark #1E293B
COLOR_MUTED = RGBColor(100, 116, 139)     # Slate Gray #64748B
COLOR_RED = RGBColor(220, 38, 38)         # Red

HEX_PRIMARY = "1E3A8A"
HEX_SECONDARY = "0F766E"
HEX_BG_LIGHT = "F8FAFC"
HEX_BG_CODE = "F1F5F9"
HEX_BORDER = "CBD5E1"
HEX_HEADER_BG = "1E3A8A"
HEX_CALLOUT_BG = "EFF6FF"
HEX_CALLOUT_BORDER = "3B82F6"
HEX_WARNING_BG = "FEF2F2"
HEX_WARNING_BORDER = "EF4444"
HEX_CONFIRMED = "DBEAFE"


def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def set_cell_background(cell, hex_color):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shd)


def set_table_borders(table, hex_color=HEX_BORDER):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="{hex_color}"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{hex_color}"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def add_callout_box(doc, text, title="LƯU Ý QUAN TRỌNG", is_warning=False):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    
    bg_color = HEX_WARNING_BG if is_warning else HEX_CALLOUT_BG
    border_color = HEX_WARNING_BORDER if is_warning else HEX_CALLOUT_BORDER
    
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'  <w:top w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    cell._tc.get_or_add_tcPr().append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run_title = p.add_run(f"{title}: ")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(10.5)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(220, 38, 38) if is_warning else COLOR_PRIMARY
    
    run_text = p.add_run(text)
    run_text.font.name = "Arial"
    run_text.font.size = Pt(10)
    run_text.font.color.rgb = COLOR_DARK
    
    doc.add_paragraph()


def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    
    set_cell_background(cell, HEX_BG_CODE)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="16" w:space="0" w:color="{HEX_PRIMARY}"/>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'</w:tcBorders>'
    )
    cell._tc.get_or_add_tcPr().append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(13)
    
    run = p.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(30, 41, 59)
    
    doc.add_paragraph()


def build_document():
    doc = Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    # --- HEADER TIÊU ĐỀ ---
    p_header = doc.add_paragraph()
    p_header.paragraph_format.space_after = Pt(2)
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_uni = p_header.add_run("ĐẠI HỌC BÁCH KHOA HÀ NỘI — TRƯỜNG CÔNG NGHỆ THÔNG TIN VÀ TRUYỀN THÔNG\n")
    r_uni.font.name = "Arial"
    r_uni.font.size = Pt(10)
    r_uni.font.bold = True
    r_uni.font.color.rgb = COLOR_MUTED
    
    r_sub = p_header.add_run("ĐỒ ÁN TỐT NGHIỆP: HỆ THỐNG PHÁT HIỆN BIỂN QUẢNG CÁO LẤN CHIẾM VỈA HÈ QUA CCTV\n")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(11)
    r_sub.font.bold = True
    r_sub.font.color.rgb = COLOR_SECONDARY
    
    r_badge = p_header.add_run("CHƯƠNG TRÌNH ĐÀO TẠO NỘI BỘ — HỌC HIỂU ĐỒ ÁN TỪ CON SỐ 0")
    r_badge.font.name = "Arial"
    r_badge.font.size = Pt(9.5)
    r_badge.font.italic = True
    r_badge.font.color.rgb = COLOR_MUTED

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(4)
    p_div.paragraph_format.space_after = Pt(14)
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_div = p_div.add_run("━" * 50)
    r_div.font.color.rgb = COLOR_PRIMARY

    # Tiêu đề bài học
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_after = Pt(8)
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_b7 = p_title.add_run("BÀI 7: MÁY TRẠNG THÁI & BỘ LỌC CHUỖI THỜI GIAN (TEMPORAL VERIFIER)\n")
    r_b7.font.name = "Arial"
    r_b7.font.size = Pt(18)
    r_b7.font.bold = True
    r_b7.font.color.rgb = COLOR_PRIMARY
    
    r_b7_sub = p_title.add_run("Bộ lọc thông minh 3 cấp độ: Loại bỏ người vác biển đi qua, khử nhấp nháy và chống báo giả 97.7%")
    r_b7_sub.font.name = "Arial"
    r_b7_sub.font.size = Pt(13)
    r_b7_sub.font.bold = True
    r_b7_sub.font.color.rgb = COLOR_SECONDARY

    # Bảng siêu dữ liệu
    tbl_meta = doc.add_table(rows=3, cols=2)
    tbl_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Tệp mã nguồn tương ứng:", "tracking/verifier.py (Tham chiếu: reports/benchmark_report.md)"),
        ("Cấp độ bài học:", "Giai đoạn 4 — Temporal Reasoning & False Alarm Elimination"),
        ("Đối tượng học tập:", "Người học từ con số 0, cần hiểu cách hệ thống chuyển trạng thái và loại trừ người đi bộ vác biển")
    ]
    for i, (k, v) in enumerate(meta_data):
        c0 = tbl_meta.cell(i, 0)
        c1 = tbl_meta.cell(i, 1)
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        set_cell_background(c0, HEX_BG_LIGHT)
        set_cell_background(c1, HEX_BG_LIGHT)
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        r0.font.name = "Arial"
        r0.font.size = Pt(9.5)
        r0.font.bold = True
        r0.font.color.rgb = COLOR_DARK
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v)
        r1.font.name = "Arial"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = COLOR_DARK
    set_table_borders(tbl_meta, HEX_BORDER)
    doc.add_paragraph()

    # --- PHẦN 1: HÔM NAY HỌC GÌ? ---
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(4)
    r_h1 = h1.add_run("PHẦN 1: HÔM NAY HỌC GÌ?")
    r_h1.font.name = "Arial"
    r_h1.font.size = Pt(13)
    r_h1.font.bold = True
    r_h1.font.color.rgb = COLOR_PRIMARY

    p_p1 = doc.add_paragraph()
    p_p1.paragraph_format.line_spacing = Pt(14)
    p_p1.paragraph_format.space_after = Pt(6)
    p_p1.add_run(
        "• Tên bài học: Máy trạng thái hữu hạn (Finite State Machine) và Bộ lọc xác thực chuỗi thời gian trong tracking/verifier.py.\n"
        "• Mục tiêu cốt lõi:\n"
        "   1. Hiểu vòng đời 3 trạng thái của một vật thể biển hiệu: NORMAL (Hợp lệ) -> SUSPECTED (Nghi vấn) -> CONFIRMED (Xác nhận vi phạm).\n"
        "   2. Nắm vững cơ chế Cửa sổ trượt kiểm tra tính đứng yên (Sliding Window Stationary Check <= 50px): làm sao phân biệt biển đặt cố định với người đi bộ vác biển đi ngang qua.\n"
        "   3. Hiểu cơ chế Khử trùng lặp không gian (Spatial Deduplication) và Cooldown 60s: giải quyết triệt để vấn đề nhảy ID khi xe cộ che khuất.\n"
        "• Vì sao quan trọng đối với đồ án?\n"
        "   Trong nghiên cứu bóc tách (Ablation Study) tại Chương 7 báo cáo tốt nghiệp, module TemporalVerifier giúp GIẢM ĐẾN 97.7% BÁO ĐỘNG GIẢ (từ 87 cảnh báo rác xuống còn đúng 2 sự kiện vi phạm thực tế)! Đây là chốt chặn thông minh nhất của toàn hệ thống."
    )

    # --- PHẦN 2: KIẾN THỨC NỀN TẢNG ---
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(4)
    r_h2 = h2.add_run("PHẦN 2: KIẾN THỨC NỀN TẢNG")
    r_h2.font.name = "Arial"
    r_h2.font.size = Pt(13)
    r_h2.font.bold = True
    r_h2.font.color.rgb = COLOR_PRIMARY

    p_p2 = doc.add_paragraph()
    p_p2.paragraph_format.line_spacing = Pt(14)
    p_p2.paragraph_format.space_after = Pt(6)
    p_p2.add_run(
        "1. Máy trạng thái 3 cấp độ (TrackState):\n"
        "   • NORMAL: Biển hiệu không lấn chiếm vỉa hè (hoặc đã rời khỏi vỉa hè). Màu hiển thị: Xanh lá cây.\n"
        "   • SUSPECTED (Nghi vấn vi phạm): Biển thỏa mãn điều kiện hình học nhưng CHƯA ĐỦ ĐIỀU KIỆN PHẠT (chưa đủ 15 frames, hoặc đang di chuyển liên tục, hoặc conf trung bình quá thấp). Màu hiển thị: Màu Vàng/Cam. Ở trạng thái này, hệ thống CHƯA LẬP HỒ SƠ PHẠT!\n"
        "   • CONFIRMED (Đã xác nhận vi phạm): Biển duy trì đứng yên trên vỉa hè >= 15 frames (~0.5s - 3s). Màu hiển thị: Đỏ rực. Chỉ khi đạt CONFIRMED, hệ thống mới chụp ảnh bằng chứng và lưu vào CSDL SQLite!\n\n"
        "2. Bộ lọc vật thể đứng yên (Sliding Window Stationary Check):\n"
        "   • Cửa sổ trượt 1 giây gần nhất: Lưu vết tọa độ tâm (cx, cy) của vật thể.\n"
        "   • Tính độ dời lớn nhất giữa 2 thời điểm bất kỳ trong cửa sổ:\n"
        "     Nếu max_d <= 50 pixels -> Vật thể ĐỨNG YÊN (biển đặt cố định).\n"
        "     Nếu max_d > 50 pixels -> Vật thể ĐANG DI CHUYỂN (người đang vác biển đi ngang qua vỉa hè) -> Lập tức giữ ở SUSPECTED, không phạt!\n\n"
        "3. Khử trùng lặp không gian & Thời gian (Spatial & Temporal Deduplication):\n"
        "   • Cooldown Timer (60 giây): Một biển khi đã bị phạt thì trong 60 giây tiếp theo, dù nó vẫn đứng đó, hệ thống không gửi thêm cảnh báo nào nữa.\n"
        "   • Bán kính không gian (Spatial Radius R <= 50px): Nếu camera bị rung lắc hoặc có xe buýt che khuất khiến ByteTrack mất dấu và cấp một ID mới (ví dụ ID #8 thay cho ID #1) tại đúng vị trí cũ, hệ thống nhận diện đây là cùng một tọa độ và áp dụng Cooldown, ngăn chặn 100% việc lập biên bản lần 2!"
    )

    add_callout_box(
        doc,
        "Ablation Study (Nghiên cứu bóc tách): Nếu chạy không có TemporalVerifier (đánh giá từng frame độc lập), "
        "hệ thống bắn ra 87 cảnh báo/phút (spam liên tục). Khi bật TemporalVerifier với 15 frames và Cooldown 60s, "
        "số cảnh báo giảm xuống còn đúng 2 sự kiện chính xác, triệt tiêu 97.7% báo rác!",
        "KẾT QUẢ KHOA HỌC CHỨNG MINH HIỆU QUẢ",
        is_warning=False
    )

    # --- PHẦN 3: VÍ DỤ ĐỜI SỐNG ---
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(12)
    h3.paragraph_format.space_after = Pt(4)
    r_h3 = h3.add_run("PHẦN 3: VÍ DỤ ĐỜI SỐNG MINH HỌA")
    r_h3.font.name = "Arial"
    r_h3.font.size = Pt(13)
    r_h3.font.bold = True
    r_h3.font.color.rgb = COLOR_PRIMARY

    p_p3 = doc.add_paragraph()
    p_p3.paragraph_format.line_spacing = Pt(14)
    p_p3.paragraph_format.space_after = Pt(6)
    p_p3.add_run(
        "Hãy tưởng tượng quy định xử phạt dừng đỗ xe của Cảnh sát giao thông:\n"
        "• Tình huống 1: Một chiếc ô tô dừng đèn đỏ mất 15 giây, hoặc một người khuân vác chiếc biển quảng cáo đi bộ qua vỉa hè để mang vào trong nhà $\\rightarrow$ Công an có chạy ra lập biên bản phạt lỗi 'lấn chiếm lòng đường' không? Không! Vì người ta đang di chuyển hoặc tạm thời dừng lại, hành vi chưa cố định!\n"
        "• Tình huống 2: Người chủ quán bê chiếc biển đặt cố định xuống vỉa hè, tắt đèn đi vào trong nhà, chiếc biển đứng trơ trơ suốt nhiều phút $\\rightarrow$ Đã đứng yên cố định và chiếm dụng lối đi $\\rightarrow$ Lập tức lập biên bản xử phạt!\n"
        "• Trong đồ án của chúng ta: TemporalVerifier đóng vai trò phân biệt sự khác nhau giữa 'đi ngang qua vỉa hè' và 'cố tình đặt chiếm dụng vỉa hè'!"
    )

    # --- PHẦN 4: LIÊN HỆ VỚI CODE THẬT ---
    h4 = doc.add_paragraph()
    h4.paragraph_format.space_before = Pt(12)
    h4.paragraph_format.space_after = Pt(4)
    r_h4 = h4.add_run("PHẦN 4: LIÊN HỆ VỚI MÃ NGUỒN THỰC TẾ (tracking/verifier.py)")
    r_h4.font.name = "Arial"
    r_h4.font.size = Pt(13)
    r_h4.font.bold = True
    r_h4.font.color.rgb = COLOR_PRIMARY

    p_p4 = doc.add_paragraph()
    p_p4.paragraph_format.line_spacing = Pt(14)
    p_p4.paragraph_format.space_after = Pt(4)
    p_p4.add_run("Đoạn mã quyết định chuyển trạng thái từ SUSPECTED sang CONFIRMED trong tracking/verifier.py (dòng 266-300):")

    code_verifier = (
        '# 1. Kiểm tra đã đạt ngưỡng thời gian liên tục chưa (>= 15 frames / >= confirm_seconds)\n'
        'if self._confirm_threshold_reached(record, current_timestamp):\n'
        '    # 2. Kiểm tra tính đứng yên trong cửa sổ trượt 1 giây (độ dời <= 50px)\n'
        '    is_stationary = self._is_stationary(record, current_timestamp)\n\n'
        '    if not is_stationary:\n'
        '        # Đang di chuyển liên tục -> Người bê biển đi ngang qua -> Giữ SUSPECTED!\n'
        '        record["status"] = TrackState.SUSPECTED\n'
        '        suspect_reason = "DI CHUYEN"\n'
        '    elif avg_conf < self.min_avg_confidence:\n'
        '        # Độ tin cậy trung bình quá thấp -> Giữ SUSPECTED\n'
        '        record["status"] = TrackState.SUSPECTED\n'
        '        suspect_reason = f"CONF {int(avg_conf*100)}% thấp"\n'
        '    else:\n'
        '        # THỎA MÃN TẤT CẢ -> CHÍNH THỨC XÁC NHẬN VI PHẠM!\n'
        '        record["status"] = TrackState.CONFIRMED\n'
        '        suspect_reason = ""\n\n'
        '    # 3. Khử trùng lặp không gian (Spatial Deduplication & Cooldown 60s)\n'
        '    if record["status"] == TrackState.CONFIRMED:\n'
        '        is_spatially_cooldown = any(\n'
        '            ((cx - c["cx"])**2 + (cy - c["cy"])**2)**0.5 <= self.spatial_cooldown_radius\n'
        '            and (current_timestamp - c["time"]) < self.cooldown_seconds\n'
        '            for c in self.confirmed_locations\n'
        '        )\n'
        '        time_ok = (current_timestamp - record["last_alert_time"]) >= self.cooldown_seconds\n\n'
        '        if not is_spatially_cooldown and time_ok:\n'
        '            is_new_alert = True  # KÍCH HOẠT LẬP HỒ SƠ PHẠT MỚI\n'
        '            record["last_alert_time"] = current_timestamp\n'
        '            self.confirmed_locations.append({"cx": cx, "cy": cy, "time": current_timestamp, "track_id": track_id})'
    )
    add_code_block(doc, code_verifier)

    # --- PHẦN 5: GIẢI THÍCH CODE TỪNG DÒNG ---
    h5 = doc.add_paragraph()
    h5.paragraph_format.space_before = Pt(12)
    h5.paragraph_format.space_after = Pt(4)
    r_h5 = h5.add_run("PHẦN 5: GIẢI THÍCH CODE TỪNG DÒNG VÀ Ý NGHĨA KỸ THUẬT")
    r_h5.font.name = "Arial"
    r_h5.font.size = Pt(13)
    r_h5.font.bold = True
    r_h5.font.color.rgb = COLOR_PRIMARY

    p_p5 = doc.add_paragraph()
    p_p5.paragraph_format.line_spacing = Pt(14)
    p_p5.paragraph_format.space_after = Pt(6)
    p_p5.add_run(
        "1. Hàm _is_stationary(record, now):\n"
        "• Lấy các điểm tọa độ tâm (cx, cy) trong hàng đợi pos_window ở 1 giây gần nhất.\n"
        "• Tính khoảng cách lớn nhất max_d giữa hai vị trí bất kỳ. Nếu người đi bộ vác biển đi qua, trong 1 giây người đó di chuyển một đoạn 80-150px (vượt ngưỡng 50px) -> Hàm trả về False -> Không xác nhận vi phạm!\n\n"
        "2. Kiểm tra độ tin cậy trung bình (avg_conf < min_avg_confidence):\n"
        "• Tránh các trường hợp mô hình AI bị nhấp nháy nhận nhầm chiếc áo hay bóng râm thành biển quảng cáo với điểm chập chờn 0.15 - 0.20. Đồ án yêu cầu điểm trung bình cả chuỗi phải >= 0.28 mới được CONFIRMED.\n\n"
        "3. Cơ chế Khử trùng lặp không gian (is_spatially_cooldown):\n"
        "• Công thức khoảng cách Euclid: sqrt((cx - c['cx'])^2 + (cy - c['cy'])^2) <= 50.0.\n"
        "• Khi một tấm biển bị mất dấu do xe buýt che khuất và ByteTrack cấp Track ID mới, vị trí tâm (cx, cy) của nó vẫn nằm trong bán kính 50 pixel so với vị trí cũ đã phạt. Hệ thống phát hiện ra ngay và KHÔNG sinh cảnh báo trùng lặp (is_new_alert = False)!\n\n"
        "4. Cơ chế Dung sai khi bị che khuất (tolerance_frames = 2):\n"
        "• Nếu biển bị người đi bộ che mất 1 hoặc 2 frame (is_spatial_viol = False), bộ đếm không bị xóa sạch về 0 ngay, mà chỉ trừ nhẹ (popleft). Khi người đó bước qua, bộ đếm nhanh chóng phục hồi lại."
    )

    # --- PHẦN 6: MINH HỌA QUÁ TRÌNH CHẠY ---
    h6 = doc.add_paragraph()
    h6.paragraph_format.space_before = Pt(12)
    h6.paragraph_format.space_after = Pt(4)
    r_h6 = h6.add_run("PHẦN 6: MINH HỌA VÒNG ĐỜI TRẠNG THÁI CỦA MỘT TẤM BIỂN")
    r_h6.font.name = "Arial"
    r_h6.font.size = Pt(13)
    r_h6.font.bold = True
    r_h6.font.color.rgb = COLOR_PRIMARY

    p_p6 = doc.add_paragraph()
    p_p6.paragraph_format.space_after = Pt(6)
    p_p6.add_run("Bảng mô phỏng dữ liệu trạng thái của Track #1 biến đổi theo thời gian:")

    tbl_flow = doc.add_table(rows=5, cols=5)
    tbl_flow.alignment = WD_TABLE_ALIGNMENT.CENTER
    flow_headers = ["Thời điểm", "Số frames vi phạm", "Độ dời tâm (max_d)", "Trạng thái (Status)", "Hành động hệ thống"]
    for j, h in enumerate(flow_headers):
        c = tbl_flow.cell(0, j)
        set_cell_background(c, HEX_HEADER_BG)
        set_cell_margins(c, 100, 100, 120, 120)
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    flow_rows = [
        ("Frame 1 (t = 0.03s)", "1 frame", "0 px", "SUSPECTED (Nghi vấn)", "Chưa đủ 15 frames -> Chưa phạt"),
        ("Frame 10 (t = 0.33s)", "10 frames", "3 px (đứng yên)", "SUSPECTED (Nghi vấn)", "Tiếp tục tích lũy chuỗi thời gian"),
        ("Frame 15 (t = 0.50s)", "15 frames (ĐỦ)", "4 px (<= 50px)", "CONFIRMED (XÁC NHẬN)", "KÍCH HOẠT PHẠT! Lưu ảnh & CSDL"),
        ("Frame 16-1800 (1 phút)", "Duy trì vi phạm", "5 px (đứng yên)", "CONFIRMED (Đang Cooldown)", "Đã phạt rồi -> Bỏ qua, không spam!")
    ]
    for i, row in enumerate(flow_rows):
        for j, val in enumerate(row):
            c = tbl_flow.cell(i + 1, j)
            bg = HEX_BG_LIGHT if i % 2 == 0 else "FFFFFF"
            set_cell_background(c, bg)
            set_cell_margins(c, 80, 80, 100, 100)
            p = c.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Arial"
            r.font.size = Pt(9)
            if j == 3 and "CONFIRMED (XÁC NHẬN)" in val:
                r.font.bold = True
                r.font.color.rgb = COLOR_RED
            else:
                r.font.color.rgb = COLOR_DARK
    set_table_borders(tbl_flow, HEX_BORDER)
    doc.add_paragraph()

    # --- PHẦN 7: LỖI THƯỜNG GẶP ---
    h7 = doc.add_paragraph()
    h7.paragraph_format.space_before = Pt(12)
    h7.paragraph_format.space_after = Pt(4)
    r_h7 = h7.add_run("PHẦN 7: LỖI THƯỜNG GẶP CỦA NGƯỜI MỚI BẮT ĐẦU VÀ CÁCH KHẮC PHỤC")
    r_h7.font.name = "Arial"
    r_h7.font.size = Pt(13)
    r_h7.font.bold = True
    r_h7.font.color.rgb = COLOR_PRIMARY

    p_p7 = doc.add_paragraph()
    p_p7.paragraph_format.line_spacing = Pt(14)
    p_p7.paragraph_format.space_after = Pt(6)
    p_p7.add_run(
        "1. Lỗi so sánh tâm hiện tại với tâm lúc XUẤT HIỆN (Displacement from Origin):\n"
        "• Lỗi cũ (Version 1.0): So sánh tâm frame 60 với tâm lúc frame 1. Nếu người ta bê biển từ ngoài vào rồi ĐẶT XUỐNG ĐỨNG YÊN, khoảng cách so với lúc đầu rất lớn khiến biển đứng yên mãi mãi không bao giờ được xác nhận vi phạm!\n"
        "• Khắc phục (Version 2.1): Chuyển sang Cửa sổ trượt (Sliding Window) 1 giây gần nhất. Chỉ cần trong 1 giây gần đây biển đứng yên thì lập tức được xác nhận vi phạm!\n\n"
        "2. Nhầm lẫn giữa SUSPECTED và CONFIRMED:\n"
        "• Sinh viên hay nhầm: Cứ thấy biển hiện lên màu vàng là tưởng hệ thống đã ghi nhận vi phạm vào database.\n"
        "• Thực tế: SUSPECTED chỉ là theo dõi nội bộ trong RAM, chỉ có CONFIRMED mới là hành vi vi phạm chính thức được lập hồ sơ pháp lý!"
    )

    # --- PHẦN 8: BÀI TẬP KIỂM TRA ---
    h8 = doc.add_paragraph()
    h8.paragraph_format.space_before = Pt(12)
    h8.paragraph_format.space_after = Pt(4)
    r_h8 = h8.add_run("PHẦN 8: BÀI TẬP KIỂM TRA HIỂU BÀI (TỰ ÔN TẬP)")
    r_h8.font.name = "Arial"
    r_h8.font.size = Pt(13)
    r_h8.font.bold = True
    r_h8.font.color.rgb = COLOR_PRIMARY

    p_p8 = doc.add_paragraph()
    p_p8.paragraph_format.line_spacing = Pt(14)
    p_p8.paragraph_format.space_after = Pt(10)
    p_p8.add_run(
        "Bạn hãy tự trả lời 3 câu hỏi dưới đây trước khi chuyển sang Bài học số 8:\n\n"
        "• Câu 1: Sự khác nhau bản chất giữa trạng thái SUSPECTED và trạng thái CONFIRMED là gì? Ở trạng thái nào thì hệ thống mới chụp ảnh bằng chứng và ghi hồ sơ vào CSDL?\n\n"
        "• Câu 2: Giả sử một người đi bộ vác tấm biển quảng cáo đi ngang qua vỉa hè trong 5 giây. Tại sao hệ thống KHÔNG BAO GIỜ xác nhận vi phạm (không bao giờ chuyển sang CONFIRMED) đối với người này?\n\n"
        "• Câu 3: Cơ chế Khử trùng lặp không gian (Spatial Deduplication) với bán kính R = 50px và Cooldown 60s giải quyết vấn đề gì khi camera bị xe cộ đi qua che khuất tạm thời?"
    )

    try:
        doc.save(OUTPUT_FILE)
        print(f"[THÀNH CÔNG] Đã tạo tệp tài liệu bài học: {OUTPUT_FILE}")
    except PermissionError:
        print(f"[CẢNH BÁO] Không thể ghi đè {OUTPUT_FILE} do file đang được mở trong Microsoft Word. Vui lòng đóng file Word và chạy lại.")


if __name__ == "__main__":
    build_document()
