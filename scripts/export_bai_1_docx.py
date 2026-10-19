"""
Script xuất bản bài học Bài 1 thành file Word .docx chuyên nghiệp
Lưu vào thư mục: tai_lieu_hoc_tap/Bai_01_Video_So_va_Khung_Hinh_OpenCV.docx
"""

import os
import sys
from pathlib import Path

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

ROOT_DIR = Path(__file__).resolve().parent.parent
OUTPUT_DIR = ROOT_DIR / "tai_lieu_hoc_tap"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
OUTPUT_FILE = str(OUTPUT_DIR / "Bai_01_Video_So_va_Khung_Hinh_OpenCV.docx")

# Màu sắc chủ đạo học thuật
COLOR_PRIMARY = RGBColor(30, 58, 138)     # Navy #1E3A8A
COLOR_SECONDARY = RGBColor(15, 118, 110)  # Teal #0F766E
COLOR_DARK = RGBColor(30, 41, 59)         # Slate Dark #1E293B
COLOR_MUTED = RGBColor(100, 116, 139)     # Slate Gray #64748B
COLOR_RED = RGBColor(220, 38, 38)         # Red
COLOR_GREEN = RGBColor(16, 185, 129)      # Green

HEX_PRIMARY = "1E3A8A"
HEX_SECONDARY = "0F766E"
HEX_BG_LIGHT = "F8FAFC"
HEX_BG_CODE = "F1F5F9"
HEX_BORDER = "CBD5E1"
HEX_HEADER_BG = "1E3A8A"
HEX_CALLOUT_BG = "EFF6FF"
HEX_CALLOUT_BORDER = "3B82F6"
HEX_WARNING_BG = "FEF2F2"
HEX_WARNING_BORDER = "EF4444"
HEX_CONFIRMED = "DBEAFE"


def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def set_cell_background(cell, hex_color):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shd)


def set_table_borders(table, hex_color=HEX_BORDER):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="{hex_color}"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{hex_color}"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def add_callout_box(doc, text, title="LƯU Ý QUAN TRỌNG", is_warning=False):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    
    bg_color = HEX_WARNING_BG if is_warning else HEX_CALLOUT_BG
    border_color = HEX_WARNING_BORDER if is_warning else HEX_CALLOUT_BORDER
    
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Border trái đậm
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'  <w:top w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    cell._tc.get_or_add_tcPr().append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run_title = p.add_run(f"{title}: ")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(10.5)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(220, 38, 38) if is_warning else COLOR_PRIMARY
    
    run_text = p.add_run(text)
    run_text.font.name = "Arial"
    run_text.font.size = Pt(10)
    run_text.font.color.rgb = COLOR_DARK
    
    doc.add_paragraph()


def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    
    set_cell_background(cell, HEX_BG_CODE)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="16" w:space="0" w:color="{HEX_PRIMARY}"/>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'</w:tcBorders>'
    )
    cell._tc.get_or_add_tcPr().append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(13)
    
    run = p.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(30, 41, 59)
    
    doc.add_paragraph()


def build_document():
    doc = Document()
    
    # Cấu hình lề trang tiêu chuẩn
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    # --- HEADER TIÊU ĐỀ ---
    p_header = doc.add_paragraph()
    p_header.paragraph_format.space_after = Pt(2)
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_uni = p_header.add_run("ĐẠI HỌC BÁCH KHOA HÀ NỘI — TRƯỜNG CÔNG NGHỆ THÔNG TIN VÀ TRUYỀN THÔNG\n")
    r_uni.font.name = "Arial"
    r_uni.font.size = Pt(10)
    r_uni.font.bold = True
    r_uni.font.color.rgb = COLOR_MUTED
    
    r_sub = p_header.add_run("ĐỒ ÁN TỐT NGHIỆP: HỆ THỐNG PHÁT HIỆN BIỂN QUẢNG CÁO LẤN CHIẾM VỈA HÈ QUA CCTV\n")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(11)
    r_sub.font.bold = True
    r_sub.font.color.rgb = COLOR_SECONDARY
    
    r_badge = p_header.add_run("CHƯƠNG TRÌNH ĐÀO TẠO NỘI BỘ — HỌC HIỂU ĐỒ ÁN TỪ CON SỐ 0")
    r_badge.font.name = "Arial"
    r_badge.font.size = Pt(9.5)
    r_badge.font.italic = True
    r_badge.font.color.rgb = COLOR_MUTED

    # Đường kẻ ngang phân cách
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(4)
    p_div.paragraph_format.space_after = Pt(14)
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_div = p_div.add_run("━" * 50)
    r_div.font.color.rgb = COLOR_PRIMARY

    # Tiêu đề bài học
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_after = Pt(8)
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_b1 = p_title.add_run("BÀI 1: VIDEO SỐ VÀ KHUNG HÌNH (FRAME)\n")
    r_b1.font.name = "Arial"
    r_b1.font.size = Pt(18)
    r_b1.font.bold = True
    r_b1.font.color.rgb = COLOR_PRIMARY
    
    r_b1_sub = p_title.add_run("Cách máy tính tiếp nhận và xử lý dữ liệu từ Camera CCTV qua OpenCV")
    r_b1_sub.font.name = "Arial"
    r_b1_sub.font.size = Pt(13)
    r_b1_sub.font.bold = True
    r_b1_sub.font.color.rgb = COLOR_SECONDARY

    # Bảng siêu dữ liệu bài học
    tbl_meta = doc.add_table(rows=3, cols=2)
    tbl_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Tệp mã nguồn tương ứng:", "camera/video_reader.py (Tham chiếu: camera/stream_reader.py, main.py)"),
        ("Cấp độ bài học:", "Giai đoạn 1 — Nhập môn Computer Vision & OpenCV"),
        ("Đối tượng học tập:", "Người học từ con số 0, cần hiểu bản chất để sửa code và bảo vệ đồ án")
    ]
    for i, (k, v) in enumerate(meta_data):
        c0 = tbl_meta.cell(i, 0)
        c1 = tbl_meta.cell(i, 1)
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        set_cell_background(c0, HEX_BG_LIGHT)
        set_cell_background(c1, HEX_BG_LIGHT)
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        r0.font.name = "Arial"
        r0.font.size = Pt(9.5)
        r0.font.bold = True
        r0.font.color.rgb = COLOR_DARK
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v)
        r1.font.name = "Arial"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = COLOR_DARK
    set_table_borders(tbl_meta, HEX_BORDER)
    doc.add_paragraph()

    # --- PHẦN 1: HÔM NAY HỌC GÌ? ---
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(4)
    r_h1 = h1.add_run("PHẦN 1: HÔM NAY HỌC GÌ?")
    r_h1.font.name = "Arial"
    r_h1.font.size = Pt(13)
    r_h1.font.bold = True
    r_h1.font.color.rgb = COLOR_PRIMARY

    p_p1 = doc.add_paragraph()
    p_p1.paragraph_format.line_spacing = Pt(14)
    p_p1.paragraph_format.space_after = Pt(6)
    p_p1.add_run(
        "• Tên bài học: Bản chất của Video kỹ thuật số và cơ chế bóc tách từng khung hình bằng OpenCV.\n"
        "• Mục tiêu cốt lõi:\n"
        "   1. Hiểu bản chất video dưới góc nhìn máy tính (tập hợp các bức ảnh tĩnh liên tiếp).\n"
        "   2. Nắm vững hệ trục tọa độ ảnh (Pixel Coordinates) — gốc (0, 0) ở góc trên cùng bên trái.\n"
        "   3. Hiểu từng dòng code trong module camera/video_reader.py để làm chủ dòng dữ liệu đầu vào của toàn hệ thống.\n"
        "• Vì sao quan trọng đối với đồ án?\n"
        "   Mọi mô hình trí tuệ nhân tạo (YOLO), thuật toán theo dõi (ByteTrack) hay phán quyết lấn chiếm vỉa hè đều bắt đầu từ một khung hình (frame). Nếu không hiểu cách máy tính thu nhận và biểu diễn khung hình, bạn sẽ không thể lý giải được tọa độ bounding box hay vị trí tiếp đất của biển hiệu."
    )

    # --- PHẦN 2: KIẾN THỨC NỀN TẢNG ---
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(4)
    r_h2 = h2.add_run("PHẦN 2: KIẾN THỨC NỀN TẢNG")
    r_h2.font.name = "Arial"
    r_h2.font.size = Pt(13)
    r_h2.font.bold = True
    r_h2.font.color.rgb = COLOR_PRIMARY

    p_p2 = doc.add_paragraph()
    p_p2.paragraph_format.line_spacing = Pt(14)
    p_p2.paragraph_format.space_after = Pt(6)
    p_p2.add_run(
        "1. Khung hình (Frame):\n"
        "Video thực chất không phải là một luồng chuyển động liên tục, mà là tập hợp của hàng nghìn bức ảnh tĩnh được xếp nối tiếp nhau và trình chiếu với tốc độ cực nhanh.\n\n"
        "2. Tốc độ khung hình (FPS - Frames Per Second):\n"
        "Số lượng bức ảnh tĩnh được ghi nhận hoặc trình chiếu trong 1 giây. Nếu camera giám sát đạt FPS = 25 hoặc 30, nghĩa là mỗi giây trôi qua máy tính phải xử lý 25 đến 30 bức ảnh tĩnh.\n\n"
        "3. Ma trận điểm ảnh (Pixel Matrix):\n"
        "Một bức ảnh Full HD (1920x1080) hay HD (1280x720) đối với máy tính là một mảng 3 chiều (NumPy array) gồm:\n"
        "   - Chiều cao (Height - H): Số dòng điểm ảnh (ví dụ: 720 dòng).\n"
        "   - Chiều rộng (Width - W): Số cột điểm ảnh (ví dụ: 1280 cột).\n"
        "   - Kênh màu (Channels - C): Gồm 3 giá trị màu Blue, Green, Red (chuẩn BGR của thư viện OpenCV).\n\n"
        "4. Hệ trục tọa độ ảnh (Pixel Coordinate System) — CỰC KỲ ĐẶC BIỆT:\n"
        "Khác với toán học phổ thông (gốc tọa độ ở góc dưới bên trái), trong màn hình máy tính và thị giác máy tính:\n"
        "   - Gốc (0, 0) nằm tại GÓC TRÊN CÙNG BÊN TRÁI của ảnh.\n"
        "   - Trục X tăng dần từ TRÁI SANG PHẢI (0 → Width).\n"
        "   - Trục Y tăng dần từ TRÊN XUỐNG DƯỚI (0 → Height).\n"
        "   - Hệ quả: Đỉnh của biển hiệu có tọa độ Y NHỎ HƠN, còn chân tiếp đất của biển có tọa độ Y LỚN HƠN."
    )

    add_callout_box(
        doc,
        "Quy ước hệ tọa độ ảnh là nền tảng để hiểu tại sao trong hàm phán quyết vi phạm (violation/checker.py), "
        "chân biển chạm đất lại có tọa độ y = y2 (cạnh đáy) lớn hơn đỉnh y = y1 (cạnh trên). Đừng bao giờ nhầm lẫn!",
        "ĐIỂM HỘI ĐỒNG THƯỜNG HỎI",
        is_warning=True
    )

    # --- PHẦN 3: VÍ DỤ ĐỜI SỐNG ---
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(12)
    h3.paragraph_format.space_after = Pt(4)
    r_h3 = h3.add_run("PHẦN 3: VÍ DỤ ĐỜI SỐNG MINH HỌA")
    r_h3.font.name = "Arial"
    r_h3.font.size = Pt(13)
    r_h3.font.bold = True
    r_h3.font.color.rgb = COLOR_PRIMARY

    p_p3 = doc.add_paragraph()
    p_p3.paragraph_format.line_spacing = Pt(14)
    p_p3.paragraph_format.space_after = Pt(6)
    p_p3.add_run(
        "Hãy tưởng tượng trò chơi 'Sách lật hoạt hình' (Flipbook) thời thơ ấu:\n"
        "• Bạn vẽ một hình người que ở góc trang giấy, mỗi trang sau người que bước một bước nhỏ.\n"
        "• Khi bạn dùng ngón tay miết mép quyển sách để các trang lật thật nhanh (tầm 25-30 trang/giây), mắt người sẽ thấy người que đang đi bộ mượt mà.\n"
        "• Trong đồ án của chúng ta: Camera CCTV chính là người đang miết lật từng trang sách, còn chương trình Python của chúng ta là người giám sát, mở từng trang giấy ra, soi xem trên trang giấy đó có biển hiệu nào đặt lên vỉa hè hay không."
    )

    # --- PHẦN 4: LIÊN HỆ VỚI CODE THẬT ---
    h4 = doc.add_paragraph()
    h4.paragraph_format.space_before = Pt(12)
    h4.paragraph_format.space_after = Pt(4)
    r_h4 = h4.add_run("PHẦN 4: LIÊN HỆ VỚI MÃ NGUỒN THỰC TẾ (camera/video_reader.py)")
    r_h4.font.name = "Arial"
    r_h4.font.size = Pt(13)
    r_h4.font.bold = True
    r_h4.font.color.rgb = COLOR_PRIMARY

    p_p4 = doc.add_paragraph()
    p_p4.paragraph_format.line_spacing = Pt(14)
    p_p4.paragraph_format.space_after = Pt(4)
    p_p4.add_run("Toàn bộ mã nguồn cốt lõi trong tệp camera/video_reader.py được tổ chức ngắn gọn như sau:")

    code_content = (
        'import os\n'
        'import cv2\n\n'
        'class VideoReader:\n'
        '    def __init__(self, source_path: str):\n'
        '        self.source_path = source_path\n'
        '        self.cap = cv2.VideoCapture(source_path)\n'
        '        if not self.cap.isOpened():\n'
        '            raise ValueError(f"Không thể mở nguồn video: {source_path}")\n'
        '            \n'
        '        self.width = int(self.cap.get(cv2.CAP_PROP_FRAME_WIDTH))\n'
        '        self.height = int(self.cap.get(cv2.CAP_PROP_FRAME_HEIGHT))\n'
        '        self.fps = self.cap.get(cv2.CAP_PROP_FPS) or 30.0\n'
        '        self.total_frames = int(self.cap.get(cv2.CAP_PROP_FRAME_COUNT))\n\n'
        '    def read(self):\n'
        '        """Đọc 1 frame tiếp theo (ret, frame)"""\n'
        '        return self.cap.read()\n\n'
        '    def reset(self):\n'
        '        """Tua lại đầu video"""\n'
        '        self.cap.set(cv2.CAP_PROP_POS_FRAMES, 0)\n\n'
        '    def release(self):\n'
        '        """Giải phóng bộ nhớ camera"""\n'
        '        if self.cap and self.cap.isOpened():\n'
        '            self.cap.release()'
    )
    add_code_block(doc, code_content)

    # --- PHẦN 5: GIẢI THÍCH CODE TỪNG DÒNG ---
    h5 = doc.add_paragraph()
    h5.paragraph_format.space_before = Pt(12)
    h5.paragraph_format.space_after = Pt(4)
    r_h5 = h5.add_run("PHẦN 5: GIẢI THÍCH CODE TỪNG DÒNG VÀ Ý NGHĨA KỸ THUẬT")
    r_h5.font.name = "Arial"
    r_h5.font.size = Pt(13)
    r_h5.font.bold = True
    r_h5.font.color.rgb = COLOR_PRIMARY

    p_p5 = doc.add_paragraph()
    p_p5.paragraph_format.line_spacing = Pt(14)
    p_p5.paragraph_format.space_after = Pt(6)
    p_p5.add_run(
        "1. Khởi tạo đối tượng đọc video (__init__):\n"
        "• class VideoReader: Tạo ra một lớp (Class) đóng gói toàn bộ logic đọc video. Các file khác như main.py chỉ cần gọi đối tượng này mà không cần viết lại mã lặp đi lặp lại.\n"
        "• self.cap = cv2.VideoCapture(source_path): Hàm của thư viện OpenCV đảm nhiệm việc kết nối tới tệp video (hoặc luồng RTSP / Webcam) qua đường dẫn source_path.\n"
        "• if not self.cap.isOpened(): Kiểm tra xem tệp video có thực sự tồn tại hoặc luồng camera có hợp lệ hay không. Nếu đường dẫn sai, chương trình chủ động văng lỗi ValueError để lập trình viên biết chính xác nguyên nhân, tránh bị lỗi ngầm.\n"
        "• self.width và self.height: Trích xuất độ phân giải gốc của video (ví dụ: 1280x720). Giá trị này rất quan trọng để tầng ROI tự động co giãn tọa độ vỉa hè theo tỷ lệ tương ứng.\n"
        "• self.fps = self.cap.get(cv2.CAP_PROP_FPS) or 30.0: Lấy tốc độ khung hình. Nếu video bị thiếu metadata không báo FPS, tự động gán giá trị dự phòng là 30.0 FPS.\n\n"
        "2. Bóc tách từng khung hình (Hàm read):\n"
        "• return self.cap.read(): ĐÂY LÀ DÒNG CODE QUAN TRỌNG NHẤT CỦA MODULE. Mỗi lần gọi, OpenCV sẽ đọc bức ảnh tiếp theo và trả về cặp giá trị (ret, frame):\n"
        "   - ret (Boolean: True/False): Báo trạng thái đọc thành công hay thất bại (khi video kết thúc).\n"
        "   - frame (NumPy Array): Bức ảnh ma trận chứa dữ liệu pixel thực tế để nạp vào mô hình AI YOLO.\n\n"
        "3. Dọn dẹp tài nguyên (Hàm release):\n"
        "• self.cap.release(): Giải phóng bộ nhớ RAM và đóng kết nối tới camera. Nếu không gọi hàm này, hệ điều hành sẽ tiếp tục giữ cổng camera và ứng dụng khác mở camera sẽ bị báo lỗi 'Thiết bị đang bận'."
    )

    # --- PHẦN 6: MINH HỌA QUÁ TRÌNH CHẠY ---
    h6 = doc.add_paragraph()
    h6.paragraph_format.space_before = Pt(12)
    h6.paragraph_format.space_after = Pt(4)
    r_h6 = h6.add_run("PHẦN 6: MINH HỌA QUÁ TRÌNH DỮ LIỆU BIẾN ĐỔI KHI HỆ THỐNG CHẠY")
    r_h6.font.name = "Arial"
    r_h6.font.size = Pt(13)
    r_h6.font.bold = True
    r_h6.font.color.rgb = COLOR_PRIMARY

    p_p6 = doc.add_paragraph()
    p_p6.paragraph_format.space_after = Pt(6)
    p_p6.add_run("Bảng mô phỏng luồng dữ liệu khi vòng lặp giám sát của main.py gọi VideoReader:")

    tbl_flow = doc.add_table(rows=5, cols=4)
    tbl_flow.alignment = WD_TABLE_ALIGNMENT.CENTER
    flow_headers = ["Thời điểm", "Lệnh gọi", "Kết quả trả về (ret, frame)", "Hành động của hệ thống"]
    for j, h in enumerate(flow_headers):
        c = tbl_flow.cell(0, j)
        set_cell_background(c, HEX_HEADER_BG)
        set_cell_margins(c, 100, 100, 120, 120)
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    flow_rows = [
        ("Vòng 1 (t = 0.00s)", "ret, frame = reader.read()", "ret=True\nframe=(720, 1280, 3)", "Nạp ảnh 1 vào YOLO phát hiện biển hiệu"),
        ("Vòng 2 (t = 0.03s)", "ret, frame = reader.read()", "ret=True\nframe=(720, 1280, 3)", "Nạp ảnh 2 vào YOLO + ByteTrack gán ID"),
        ("Vòng N (t = 10.0s)", "ret, frame = reader.read()", "ret=True\nframe=(720, 1280, 3)", "Tiếp tục quét chuỗi thời gian"),
        ("Hết video", "ret, frame = reader.read()", "ret=False\nframe=None", "Nhận biết video kết thúc, thoát hoặc tua lại")
    ]
    for i, row in enumerate(flow_rows):
        for j, val in enumerate(row):
            c = tbl_flow.cell(i + 1, j)
            bg = HEX_BG_LIGHT if i % 2 == 0 else "FFFFFF"
            set_cell_background(c, bg)
            set_cell_margins(c, 80, 80, 100, 100)
            p = c.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Arial"
            r.font.size = Pt(9)
            r.font.color.rgb = COLOR_DARK
    set_table_borders(tbl_flow, HEX_BORDER)
    doc.add_paragraph()

    # --- PHẦN 7: LỖI THƯỜNG GẶP ---
    h7 = doc.add_paragraph()
    h7.paragraph_format.space_before = Pt(12)
    h7.paragraph_format.space_after = Pt(4)
    r_h7 = h7.add_run("PHẦN 7: LỖI THƯỜNG GẶP CỦA NGƯỜI MỚI BẮT ĐẦU VÀ CÁCH KHẮC PHỤC")
    r_h7.font.name = "Arial"
    r_h7.font.size = Pt(13)
    r_h7.font.bold = True
    r_h7.font.color.rgb = COLOR_PRIMARY

    p_p7 = doc.add_paragraph()
    p_p7.paragraph_format.line_spacing = Pt(14)
    p_p7.paragraph_format.space_after = Pt(6)
    p_p7.add_run(
        "1. Lỗi 'AttributeError: 'NoneType' object has no attribute 'shape':\n"
        "• Nguyên nhân: Khi video đọc hết khung hình cuối cùng, reader.read() trả về frame = None. Người mới thường quên kiểm tra ret mà mang ngay biến frame đi vẽ hoặc đưa vào AI, dẫn đến sập chương trình.\n"
        "• Cách viết an toàn:\n"
        "   ret, frame = reader.read()\n"
        "   if not ret:\n"
        "       break  # Dừng vòng lặp kịp thời khi hết video\n\n"
        "2. Nhầm lẫn tọa độ chiều cao màn hình:\n"
        "• Nhầm lẫn: Tưởng chân cột biển hiệu tiếp đất có y nhỏ hơn đỉnh biển hiệu.\n"
        "• Thực tế: Vì trục Y hướng xuống dưới nên chân tiếp đất luôn có y LỚN HƠN đỉnh biển hiệu (y_chân > y_đỉnh)."
    )

    # --- PHẦN 8: BÀI TẬP KIỂM TRA ---
    h8 = doc.add_paragraph()
    h8.paragraph_format.space_before = Pt(12)
    h8.paragraph_format.space_after = Pt(4)
    r_h8 = h8.add_run("PHẦN 8: BÀI TẬP KIỂM TRA HIỂU BÀI (TỰ ÔN TẬP)")
    r_h8.font.name = "Arial"
    r_h8.font.size = Pt(13)
    r_h8.font.bold = True
    r_h8.font.color.rgb = COLOR_PRIMARY

    p_p8 = doc.add_paragraph()
    p_p8.paragraph_format.line_spacing = Pt(14)
    p_p8.paragraph_format.space_after = Pt(10)
    p_p8.add_run(
        "Bạn hãy tự trả lời 3 câu hỏi dưới đây trước khi chuyển sang Bài học số 2:\n\n"
        "• Câu 1: Trong camera giám sát có kích thước 1280x720, gốc tọa độ (0, 0) nằm ở đâu? Điểm chân biển hiệu tiếp đất ở mép dưới cùng màn hình có tọa độ Y xấp xỉ bằng 0 hay xấp xỉ bằng 720?\n\n"
        "• Câu 2: Hàm reader.read() trả về hai giá trị (ret, frame). Biến ret mang kiểu dữ liệu gì (chữ, số hay boolean đúng/sai), và vai trò của nó là để báo cho hệ thống biết điều gì?\n\n"
        "• Câu 3: Nếu camera CCTV của bạn quay với tốc độ FPS = 25, điều đó có nghĩa là trong 1 phút ghi hình, hệ thống sẽ gọi hàm reader.read() để bóc tách bao nhiêu bức ảnh tĩnh (khung hình)?"
    )

    # --- PHỤ LỤC: LỘ TRÌNH TOÀN DIỆN ---
    doc.add_page_break()
    h_road = doc.add_paragraph()
    h_road.paragraph_format.space_before = Pt(10)
    h_road.paragraph_format.space_after = Pt(4)
    r_road = h_road.add_run("PHỤ LỤC: LỘ TRÌNH 10 BÀI HỌC TOÀN DIỆN CỦA ĐỒ ÁN")
    r_road.font.name = "Arial"
    r_road.font.size = Pt(14)
    r_road.font.bold = True
    r_road.font.color.rgb = COLOR_PRIMARY

    p_road_desc = doc.add_paragraph()
    p_road_desc.paragraph_format.space_after = Pt(8)
    p_road_desc.add_run("Lộ trình được thiết kế theo phương pháp Bottom-Up giúp người học nắm chắc từ gốc rễ:")

    tbl_road = doc.add_table(rows=11, cols=3)
    tbl_road.alignment = WD_TABLE_ALIGNMENT.CENTER
    road_headers = ["Bài học", "Chủ đề cốt lõi", "Tệp mã nguồn tương ứng"]
    for j, h in enumerate(road_headers):
        c = tbl_road.cell(0, j)
        set_cell_background(c, HEX_HEADER_BG)
        set_cell_margins(c, 100, 100, 120, 120)
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    lessons = [
        ("Bài 1 (Hiện tại)", "Video số & Khung hình: Cách máy tính đọc dữ liệu camera", "camera/video_reader.py"),
        ("Bài 2", "Vẽ đồ họa & Bắt sự kiện chuột: Tạo đa giác vỉa hè", "scripts/demos/mouse_events.py"),
        ("Bài 3", "Mô hình AI YOLOv8n & Bounding Box trích xuất vật thể", "detection/detector.py"),
        ("Bài 4", "Đa giác vỉa hè (ROI) & Thuật toán kiểm tra điểm tiếp đất", "roi/roi_manager.py, scripts/demos/check_violation_spatial.py"),
        ("Bài 5", "Phán quyết hình học kép: Overlap Ratio 30% bằng Shapely", "violation/checker.py"),
        ("Bài 6", "Theo dõi đa vật thể ByteTrack & Cấp phát Track ID", "tracking/tracker.py, scripts/demos/demo_tracking.py"),
        ("Bài 7", "Bộ lọc chuỗi thời gian & Máy trạng thái loại trừ báo giả", "tracking/verifier.py"),
        ("Bài 8", "Lưu bằng chứng pháp lý, Làm mờ mặt & Cơ sở dữ liệu SQLite", "evidence/saver.py, database/db.py"),
        ("Bài 9", "Ghép nối toàn chuỗi, Đa luồng RTSP & Dashboard quản trị", "main.py, camera/stream_reader.py, dashboard/web_app.py"),
        ("Bài 10", "Thẩm định học thuật, Số liệu thực nghiệm & Bảo vệ Hội đồng", "reports/benchmark_report.md, AI_PROJECT_OVERVIEW.md")
    ]
    for i, (b, t, f) in enumerate(lessons):
        for j, val in enumerate([b, t, f]):
            c = tbl_road.cell(i + 1, j)
            bg = HEX_CONFIRMED if i == 0 else (HEX_BG_LIGHT if i % 2 == 1 else "FFFFFF")
            set_cell_background(c, bg)
            set_cell_margins(c, 80, 80, 100, 100)
            p = c.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Arial"
            r.font.size = Pt(9)
            if i == 0:
                r.font.bold = True
                r.font.color.rgb = COLOR_PRIMARY
            else:
                r.font.color.rgb = COLOR_DARK
    set_table_borders(tbl_road, HEX_BORDER)

    doc.save(OUTPUT_FILE)
    print(f"[THÀNH CÔNG] Đã tạo tệp tài liệu bài học: {OUTPUT_FILE}")


if __name__ == "__main__":
    build_document()
