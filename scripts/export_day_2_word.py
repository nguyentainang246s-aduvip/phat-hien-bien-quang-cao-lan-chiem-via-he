"""
Script xuất bản tài liệu Ngày 02 thành file Word .docx chuyên nghiệp
Lưu vào thư mục: tai_lieu_hoc_tap/Bai_02_Do_Toc_Do_FPS_va_Thuat_Toan_EMA.docx
"""
import os
import sys
from pathlib import Path

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

ROOT_DIR = Path(__file__).resolve().parent.parent
OUTPUT_DIR = ROOT_DIR / "tai_lieu_hoc_tap"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
OUTPUT_FILE = str(OUTPUT_DIR / "Bai_02_Do_Toc_Do_FPS_va_Thuat_Toan_EMA.docx")

COLOR_PRIMARY = RGBColor(30, 58, 138)     # Navy #1E3A8A
COLOR_SECONDARY = RGBColor(15, 118, 110)  # Teal #0F766E
COLOR_DARK = RGBColor(30, 41, 59)         # Slate Dark #1E293B
COLOR_MUTED = RGBColor(100, 116, 139)     # Slate Gray #64748B

HEX_PRIMARY = "1E3A8A"
HEX_CALLOUT_BG = "EFF6FF"
HEX_CALLOUT_BORDER = "3B82F6"
HEX_BG_CODE = "F1F5F9"

def create_day_2_doc():
    doc = Document()
    
    # Header
    title = doc.add_heading("BÀI HỌC NGÀY 02: ĐO TỐC ĐỘ XỬ LÝ (FPS) & GIẢI THUẬT LÀM MƯỢT EMA", level=0)
    title.runs[0].font.color.rgb = COLOR_PRIMARY
    title.runs[0].font.size = Pt(20)
    
    p_meta = doc.add_paragraph("Đồ án Tốt nghiệp: Hệ thống Giám sát Biển quảng cáo Lấn chiếm Vỉa hè qua CCTV (DATN20224083)")
    p_meta.runs[0].font.italic = True
    p_meta.runs[0].font.color.rgb = COLOR_MUTED
    
    # 1. Mục tiêu
    h1 = doc.add_heading("1. Mục tiêu bài học", level=1)
    h1.runs[0].font.color.rgb = COLOR_SECONDARY
    doc.add_paragraph("• Hiểu FPS (Frames Per Second) là gì và tại sao bắt buộc phải đo trong đồ án.")
    doc.add_paragraph("• Nắm vững hiện tượng giật số FPS tức thời và cách giải quyết bằng thuật toán EMA.")
    doc.add_paragraph("• Mổ xẻ chi tiết class FPSTracker trong utils/fps_tracker.py và scripts/demos/read_video_fps.py.")
    doc.add_paragraph("• Biết cách vẽ huy hiệu FPS hiển thị chuyên nghiệp lên video.")
    
    # 2. Khái niệm FPS và EMA
    h2 = doc.add_heading("2. Bản chất Toán học: Công thức tính FPS và EMA", level=1)
    h2.runs[0].font.color.rgb = COLOR_SECONDARY
    doc.add_paragraph("FPS tức thời được tính bằng nghịch đảo thời gian xử lý 1 frame:")
    doc.add_paragraph("   FPS_tức_thời = 1.0 / delta_t (với delta_t = curr_time - prev_time)")
    doc.add_paragraph("Công thức làm mượt hàm mũ EMA (Exponential Moving Average):")
    doc.add_paragraph("   FPS_mới = alpha * FPS_tức_thời + (1 - alpha) * FPS_cũ (với alpha = 0.15)")
    
    # 3. Mổ xẻ Code
    h3 = doc.add_heading("3. Mã nguồn thực tế: Class FPSTracker (utils/fps_tracker.py)", level=1)
    h3.runs[0].font.color.rgb = COLOR_SECONDARY
    
    code = (
        "class FPSTracker:\n"
        "    def __init__(self, alpha: float = 0.15):\n"
        "        self.alpha = alpha\n"
        "        self.prev_time = None\n"
        "        self.smoothed_fps = 0.0\n\n"
        "    def update(self) -> float:\n"
        "        curr_time = time.perf_counter()\n"
        "        if self.prev_time is None:\n"
        "            self.prev_time = curr_time\n"
        "            return 0.0\n"
        "        delta_t = curr_time - self.prev_time\n"
        "        self.prev_time = curr_time\n"
        "        if delta_t > 0:\n"
        "            instant_fps = 1.0 / delta_t\n"
        "            if self.smoothed_fps == 0.0:\n"
        "                self.smoothed_fps = instant_fps\n"
        "            else:\n"
        "                self.smoothed_fps = self.alpha * instant_fps + (1.0 - self.alpha) * self.smoothed_fps\n"
        "        return self.smoothed_fps\n"
    )
    p_code = doc.add_paragraph(code)
    p_code.style.font.name = 'Consolas'
    p_code.style.font.size = Pt(9.5)
    
    # 4. Câu hỏi ôn tập
    h4 = doc.add_heading("4. Câu hỏi ôn tập bảo vệ đồ án", level=1)
    h4.runs[0].font.color.rgb = COLOR_SECONDARY
    doc.add_paragraph("1. Tại sao không dùng time.time() mà phải dùng time.perf_counter() để bấm giờ?")
    doc.add_paragraph("2. Hệ số alpha trong EMA có ý nghĩa gì? Nếu chọn alpha = 1.0 thì điều gì xảy ra?")
    doc.add_paragraph("3. Vì sao phải vẽ khung nền màu đen phía sau thông số FPS?")
    
    doc.save(OUTPUT_FILE)
    print(f"[THÀNH CÔNG] Đã tạo tệp: {OUTPUT_FILE}")

if __name__ == "__main__":
    create_day_2_doc()
